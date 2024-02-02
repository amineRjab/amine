from docx import Document

data = {
    '01/02/23': {
        '1': '18:17:36-18:53:03', '2': '19:00:11-19:36:32', '3': '19:43:44-20:20:07',
        '4': '20:29:53-21:07:51', '5': '21:17:37-21:53:08', '6': '21:58:38-22:33:08',
        '7': '22:43:04-23:14:37', '8': '23:19:26-23:57:47'
    },
    '02/02/23': {
        '9': '00:05:00-00:47:29', '10': '00:51:30-01:42:29', '11': '01:50:21-02:27:19',
        '12': '02:34:40-03:09:40', '14': '05:36:18-06:13:18', '15': '06:18:42-06:55:10',
        '16': '07:03:46-07:39:04', '17': '07:47:50-08:22:08', '18': '08:29:13-09:06:25',
        '19': '09:11:25-09:48:54', '20': '09:56:04-10:33:53'
    }
}

document = Document()

# Add table headers
table = document.add_table(rows=1, cols=3)
table.style = 'Table Grid'
table.cell(0, 0).text = 'Date'
table.cell(0, 1).text = 'Session'
table.cell(0, 2).text = 'Time Range'

# Add data rows to the table
for date, sessions in data.items():
    for session, time_range in sessions.items():
        row = table.add_row().cells
        row[0].text = date
        row[1].text = session
        row[2].text = time_range

# Save the document
document.save('table.docx')
